"""Layanan Ekspor Excel Komprehensif untuk Market Size, Market Share & Brand Leaderboard.
Menghasilkan workbook .xlsx berstandar eksekutif menggunakan openpyxl.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.market_size_service import get_market_size_and_share, get_syamanah_benchmark_report


def generate_market_size_excel_bytes() -> io.BytesIO:
    """Menghasilkan file Excel (.xlsx) komprehensif berisi seluruh analitik Market Size,
    Pangsa Pasar Syamanah, Benchmark 20 Brand per Kategori, serta Rumus Omzet Perusahaan.
    """
    data = get_market_size_and_share()
    benchmarks = data.get("benchmarks", {})
    financials = data.get("financials", {})
    standing = data.get("syamanah_standing", {})
    formulas = data.get("omzet_formulas", [])

    wb = Workbook()
    # Default sheet
    ws_summary = wb.active
    ws_summary.title = "Ringkasan Pasar & KPIs"

    # Define color schemes & styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    subhead_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")  # Blue Accent
    subhead_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    syamanah_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Soft Green
    syamanah_font = Font(name="Calibri", size=10, bold=True, color="166534")

    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # ---------------------------------------------------------
    # SHEET 1: RINGKASAN PASAR & KPIS
    # ---------------------------------------------------------
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary.cell(row=2, column=2, value="SYAMANAH GROUP — MARKET SIZE & MARKET SHARE EXECUTIVE REPORT 2026").font = title_font
    ws_summary.cell(row=3, column=2, value="Analisis Potensi Pasar B2B Apparel (TAM, SAM, SOM) & Benchmark Peringkat Nasional").font = subtitle_font

    # Section 1: TAM, SAM, SOM Table
    row_idx = 5
    ws_summary.cell(row=row_idx, column=2, value="1. HIERARKI POTENSI PASAR APPAREL & MERCHANDISE NASIONAL").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    row_idx += 1

    headers_tam = ["Tingkat Pasar", "Nilai Pasar (Rupiah)", "Estimasi Entitas / Populasi", "Volume Estimasi (Pcs)", "Definisi & Ruang Lingkup Pasar"]
    for col_idx, h in enumerate(headers_tam, start=2):
        cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    row_idx += 1

    tam_rows = [
        ("TAM (Total Addressable Market)", data.get("tam", {}).get("value_rp", 0), data.get("tam", {}).get("people", 0), data.get("tam", {}).get("pcs", 0), data.get("tam", {}).get("description", "")),
        ("SAM (Serviceable Available Market)", data.get("sam", {}).get("value_rp", 0), data.get("sam", {}).get("people", 0), data.get("sam", {}).get("pcs", 0), data.get("sam", {}).get("description", "")),
        ("SOM (Serviceable Obtainable Market)", data.get("som", {}).get("value_rp", 0), data.get("som", {}).get("people", 0), data.get("som", {}).get("pcs", 0), data.get("som", {}).get("description", "")),
    ]

    for label, val_rp, people, pcs, desc in tam_rows:
        ws_summary.cell(row=row_idx, column=2, value=label).font = bold_font
        
        c_rp = ws_summary.cell(row=row_idx, column=3, value=val_rp)
        c_rp.number_format = '"Rp "#,##0'
        c_rp.alignment = Alignment(horizontal="right")
        
        c_pop = ws_summary.cell(row=row_idx, column=4, value=people)
        c_pop.number_format = '#,##0'
        c_pop.alignment = Alignment(horizontal="right")
        
        c_pcs = ws_summary.cell(row=row_idx, column=5, value=pcs)
        c_pcs.number_format = '#,##0'
        c_pcs.alignment = Alignment(horizontal="right")
        
        ws_summary.cell(row=row_idx, column=6, value=desc).alignment = Alignment(wrap_text=True)

        for c in range(2, 7):
            ws_summary.cell(row=row_idx, column=c).border = cell_border
        row_idx += 1

    # Section 2: Financial Realization & Market Share
    row_idx += 2
    ws_summary.cell(row=row_idx, column=2, value="2. PENCAPAIAN FINANSIAL & PANGSA PASAR SYAMANAH GROUP (9 BULAN AKTUAL 2026)").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    row_idx += 1

    headers_kpi = ["Metrik Kinerja", "Nilai Realisasi / Proyeksi", "Catatan Keterangan"]
    for col_idx, h in enumerate(headers_kpi, start=2):
        cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
        cell.fill = subhead_fill
        cell.font = subhead_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
    row_idx += 1

    kpi_rows = [
        ("Total Realisasi Penjualan (Jan–Sep 2026)", financials.get("total_group_9mo", 21884488835), '"Rp "#,##0', "9 Bulan Berjalan (Syamanah + Raffiz)"),
        ("Rata-Rata Omzet Bulanan", financials.get("avg_monthly_group", 2431609871), '"Rp "#,##0', "Konsistensi Kas Masuk Riil"),
        ("Proyeksi Omzet Tahunan Penuh (Annualized)", financials.get("projected_annual_group", 29179318447), '"Rp "#,##0', "Ekstrapolasi 12 Bulan Penuh"),
        ("Pangsa Pasar Syamanah vs SOM", data.get("market_share", {}).get("vs_som_pct", 6.48) / 100, "0.00%", "Target Terjangkau di Segmen Korporat/Komunitas"),
        ("Pangsa Pasar Syamanah vs SAM", data.get("market_share", {}).get("vs_sam_pct", 0.66) / 100, "0.00%", "Pasar Apparel B2B Jawa Barat & Jawa Tengah"),
        ("Pangsa Pasar Syamanah vs TAM", data.get("market_share", {}).get("vs_tam_pct", 0.08) / 100, "0.00%", "Pangsa Pasar Seluruh Indonesia"),
        ("Potensi SAM Belum Tergarap (Unaddressed SAM)", data.get("market_share", {}).get("unaddressed_sam_rp", 0), '"Rp "#,##0', "Ruang Pertumbuhan Masif yang Masih Terbuka"),
    ]

    for label, val, num_fmt, note in kpi_rows:
        ws_summary.cell(row=row_idx, column=2, value=label).font = bold_font
        cell_val = ws_summary.cell(row=row_idx, column=3, value=val)
        cell_val.number_format = num_fmt
        cell_val.alignment = Alignment(horizontal="right")
        cell_val.font = bold_font
        ws_summary.cell(row=row_idx, column=4, value=note).font = normal_font

        for c in range(2, 5):
            ws_summary.cell(row=row_idx, column=c).border = cell_border
        row_idx += 1

    # Section 3: Operational Production Capacity
    row_idx += 2
    ws_summary.cell(row=row_idx, column=2, value="3. DATA KAPASITAS PRODUKSI & STRATEGI SUPPLY CHAIN SYAMANAH").font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    row_idx += 1

    headers_cap = ["Lini Produk", "Kapasitas In-House (Mesin)", "Kebijakan Overflow / Skalabilitas", "Peran Strategis"]
    for col_idx, h in enumerate(headers_cap, start=2):
        cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
    row_idx += 1

    cap_rows = [
        ("Jersey Sublimasi Olahraga & Tim", "75.000 pcs / bulan", "Vendor Maklon CMT Luar (Buffer Tak Terbatas)", "Core Engine (Margin Maksimal In-House)"),
        ("Jaket Bomber, Parka & Varsity", "6.000 pcs / bulan", "Lini Khusus In-House (Kapasitas Terukur)", "Produk Nilai Tambah Tinggi (AOV Tinggi)"),
        ("Seragam Kemeja Kerja, Polo & Kaos", "Sesuai Kebutuhan SPK", "Mitra Vendor Maklon CMT Terpercaya", "Produk Substitusi & Kelengkapan B2B"),
    ]

    for prod, cap, policy, role in cap_rows:
        ws_summary.cell(row=row_idx, column=2, value=prod).font = bold_font
        ws_summary.cell(row=row_idx, column=3, value=cap).font = bold_font
        ws_summary.cell(row=row_idx, column=4, value=policy).font = normal_font
        ws_summary.cell(row=row_idx, column=5, value=role).font = normal_font
        for c in range(2, 6):
            ws_summary.cell(row=row_idx, column=c).border = cell_border
        row_idx += 1

    # ---------------------------------------------------------
    # HELPER FUNCTION: POPULATE BENCHMARK SHEET
    # ---------------------------------------------------------
    def populate_benchmark_sheet(ws, title_text, items, category_badge):
        ws.views.sheetView[0].showGridLines = True
        ws.cell(row=2, column=1, value=title_text).font = title_font
        ws.cell(row=3, column=1, value=f"Kategori: {category_badge} — Sumber Data: Marketplace Big Data, Portofolio & Estimasi Industri 2026").font = subtitle_font

        headers = [
            "Peringkat", "Nama Brand / Perusahaan", "Lokasi / Kota", "Estimasi Volume Bulanan",
            "Estimasi Omzet Bulanan (Rp)", "Pangsa Pasar (%)", "Model Pengadaan / Spesialisasi",
            "Kanal Penjualan Utama", "Website Resmi", "Marketplace Resmi", "Instagram Resmi"
        ]
        
        row = 5
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = cell_border
        row += 1

        for item in items:
            is_sya = item.get("is_syamanah", False)
            cur_font = syamanah_font if is_sya else normal_font
            cur_fill = syamanah_fill if is_sya else None

            # Rank
            c_rank = ws.cell(row=row, column=1, value=item.get("rank", 0))
            c_rank.alignment = Alignment(horizontal="center")
            c_rank.font = syamanah_font if is_sya else bold_font

            # Name
            c_name = ws.cell(row=row, column=2, value=item.get("shop_name", item.get("name", "")))
            c_name.font = syamanah_font if is_sya else bold_font

            # City
            c_city = ws.cell(row=row, column=3, value=item.get("city", "-"))
            c_city.font = cur_font

            # Volume
            c_vol = ws.cell(row=row, column=4, value=item.get("volume", "-"))
            c_vol.font = cur_font
            c_vol.alignment = Alignment(horizontal="right")

            # Omzet Rp
            c_rev = ws.cell(row=row, column=5, value=item.get("monthly_rev", 0))
            c_rev.number_format = '"Rp "#,##0'
            c_rev.alignment = Alignment(horizontal="right")
            c_rev.font = syamanah_font if is_sya else bold_font

            # Share Pct
            c_pct = ws.cell(row=row, column=6, value=(item.get("share_pct", 0) / 100))
            c_pct.number_format = '0.00%'
            c_pct.alignment = Alignment(horizontal="right")
            c_pct.font = syamanah_font if is_sya else bold_font

            # Model
            c_model = ws.cell(row=row, column=7, value=item.get("procurement_model", "-"))
            c_model.font = cur_font

            # Channels
            channels = ", ".join(item.get("sales_channels", [])) if item.get("sales_channels") else "-"
            c_chan = ws.cell(row=row, column=8, value=channels)
            c_chan.font = cur_font

            # Website
            web_url = item.get("store_url", item.get("web_url", ""))
            c_web = ws.cell(row=row, column=9, value=web_url if web_url else "-")
            c_web.font = cur_font

            # Marketplace
            mp_name = item.get("marketplace_name", "-")
            mp_url = item.get("marketplace_url", "")
            mp_text = f"{mp_name} ({mp_url})" if mp_url and mp_name != "-" else mp_name
            c_mp = ws.cell(row=row, column=10, value=mp_text)
            c_mp.font = cur_font

            # Instagram
            ig_handle = item.get("social_handle")
            ig_url = item.get("social_url")
            ig_text = f"{ig_handle} ({ig_url})" if ig_handle and ig_url else (ig_handle or "-")
            c_ig = ws.cell(row=row, column=11, value=ig_text)
            c_ig.font = cur_font

            for c_idx in range(1, 12):
                cell_item = ws.cell(row=row, column=c_idx)
                cell_item.border = cell_border
                if cur_fill:
                    cell_item.fill = cur_fill

            row += 1

    # ---------------------------------------------------------
    # SHEET 2: TOP 20 JERSEY SUBLIMASI
    # ---------------------------------------------------------
    ws_jersey = wb.create_sheet(title="Top 20 Jersey Sublimasi")
    jersey_items = benchmarks.get("jersey", {}).get("top_sellers", [])
    populate_benchmark_sheet(ws_jersey, "BENCHMARK TOP 20 PRODUSEN JERSEY SUBLIMASI INDONESIA 2026", jersey_items, "Jersey Olahraga & Komunitas")

    # ---------------------------------------------------------
    # SHEET 3: TOP 20 JAKET & OUTERWEAR
    # ---------------------------------------------------------
    ws_jaket = wb.create_sheet(title="Top 20 Jaket & Outerwear")
    jaket_items = benchmarks.get("jaket", {}).get("top_sellers", [])
    populate_benchmark_sheet(ws_jaket, "BENCHMARK TOP 20 PRODUSEN JAKET & OUTERWEAR INDONESIA 2026", jaket_items, "Jaket Bomber, Parka & Varsity")

    # ---------------------------------------------------------
    # SHEET 4: TOP 20 SERAGAM KERJA
    # ---------------------------------------------------------
    ws_seragam = wb.create_sheet(title="Top 20 Seragam Kerja")
    seragam_items = benchmarks.get("seragam", {}).get("top_sellers", [])
    populate_benchmark_sheet(ws_seragam, "BENCHMARK TOP 20 PRODUSEN SERAGAM KERJA & KORPORAT INDONESIA 2026", seragam_items, "Kemeja, Workwear & Polo Seragam")

    # ---------------------------------------------------------
    # SHEET 5: TOP 20 BRAND NASIONAL
    # ---------------------------------------------------------
    ws_national = wb.create_sheet(title="Top 20 Brand Nasional")
    national_items = benchmarks.get("national_top_brands", [])
    ws_national.views.sheetView[0].showGridLines = True
    ws_national.cell(row=2, column=1, value="LEADERBOARD TOP 20 BRAND APPAREL & GARMEN NASIONAL TERBESAR SE-INDONESIA").font = title_font
    ws_national.cell(row=3, column=1, value="Peringkat Gabungan Lintas Kategori (Garmen Industri, Retail Nasional & Produsen Terpadu B2B)").font = subtitle_font

    headers_nat = ["Peringkat", "Nama Brand / Pabrik", "Estimasi Omzet Bulanan (Rp)", "Kanal Penjualan Utama", "Marketplace Resmi", "Instagram Resmi"]
    for col_idx, h in enumerate(headers_nat, start=1):
        c = ws_national.cell(row=5, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = cell_border

    row_nat = 6
    for item in national_items:
        is_sya = item.get("is_syamanah", False)
        cur_font = syamanah_font if is_sya else normal_font
        cur_fill = syamanah_fill if is_sya else None

        c_rank = ws_national.cell(row=row_nat, column=1, value=item.get("rank", 0))
        c_rank.alignment = Alignment(horizontal="center")
        c_rank.font = syamanah_font if is_sya else bold_font

        c_name = ws_national.cell(row=row_nat, column=2, value=item.get("name", ""))
        c_name.font = syamanah_font if is_sya else bold_font

        c_rev = ws_national.cell(row=row_nat, column=3, value=item.get("monthly_rev", 0))
        c_rev.number_format = '"Rp "#,##0'
        c_rev.alignment = Alignment(horizontal="right")
        c_rev.font = syamanah_font if is_sya else bold_font

        channels = ", ".join(item.get("sales_channels", [])) if item.get("sales_channels") else "-"
        c_chan = ws_national.cell(row=row_nat, column=4, value=channels)
        c_chan.font = cur_font

        mp_name = item.get("marketplace_name", "-")
        mp_url = item.get("marketplace_url", "")
        c_mp = ws_national.cell(row=row_nat, column=5, value=f"{mp_name} ({mp_url})" if mp_url and mp_name != "-" else mp_name)
        c_mp.font = cur_font

        ig_handle = item.get("social_handle")
        ig_url = item.get("social_url")
        c_ig = ws_national.cell(row=row_nat, column=6, value=f"{ig_handle} ({ig_url})" if ig_handle and ig_url else (ig_handle or "-"))
        c_ig.font = cur_font

        for c_idx in range(1, 7):
            cell_item = ws_national.cell(row=row_nat, column=c_idx)
            cell_item.border = cell_border
            if cur_fill:
                cell_item.fill = cur_fill

        row_nat += 1

    # ---------------------------------------------------------
    # SHEET 6: RUMUS OMZET PERUSAHAAN
    # ---------------------------------------------------------
    ws_formula = wb.create_sheet(title="Rumus Omzet Perusahaan")
    ws_formula.views.sheetView[0].showGridLines = True
    ws_formula.cell(row=2, column=1, value="METODOLOGI & RUMUS MATEMATIKA PERKIRAAN OMZET TIAP PERUSAHAAN (TOP 20)").font = title_font
    ws_formula.cell(row=3, column=1, value="Transparansi Rumus: Volume Output × Harga Acuan SPK B2B / Laporan Publik / Kas Riil").font = subtitle_font

    headers_f = [
        "Peringkat", "Nama Perusahaan", "Kategori", "Metode Perhitungan",
        "Persamaan Rumus Matematika", "Estimasi Volume Bulanan", "Harga Acuan B2B (Rp)",
        "Estimasi Omzet Bulanan (Rp)", "Sumber Data & Audit Note"
    ]
    for col_idx, h in enumerate(headers_f, start=1):
        c = ws_formula.cell(row=5, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = cell_border

    row_f = 6
    for item in formulas:
        is_sya = "SYAMANAH" in item.get("company_name", "").upper()
        cur_font = syamanah_font if is_sya else normal_font
        cur_fill = syamanah_fill if is_sya else None

        c_rank = ws_formula.cell(row=row_f, column=1, value=item.get("rank", 0))
        c_rank.alignment = Alignment(horizontal="center")
        c_rank.font = syamanah_font if is_sya else bold_font

        c_name = ws_formula.cell(row=row_f, column=2, value=item.get("company_name", ""))
        c_name.font = syamanah_font if is_sya else bold_font

        c_cat = ws_formula.cell(row=row_f, column=3, value=item.get("category", ""))
        c_cat.font = cur_font

        c_met = ws_formula.cell(row=row_f, column=4, value=item.get("method_label", ""))
        c_met.font = cur_font

        c_eq = ws_formula.cell(row=row_f, column=5, value=item.get("math_formula", ""))
        c_eq.font = Font(name="Consolas", size=9, bold=True, color="1E3A8A" if not is_sya else "166534")

        c_vol = ws_formula.cell(row=row_f, column=6, value=item.get("volume_display", ""))
        c_vol.font = cur_font

        c_prc = ws_formula.cell(row=row_f, column=7, value=item.get("price_unit_display", ""))
        c_prc.font = cur_font

        c_rev = ws_formula.cell(row=row_f, column=8, value=item.get("monthly_revenue_rp", 0))
        c_rev.number_format = '"Rp "#,##0'
        c_rev.alignment = Alignment(horizontal="right")
        c_rev.font = syamanah_font if is_sya else bold_font

        audit_text = f"Audit: {item.get('audit_note', '-')} | Sumber: {item.get('data_source', '-')}"
        c_aud = ws_formula.cell(row=row_f, column=9, value=audit_text)
        c_aud.font = cur_font
        c_aud.alignment = Alignment(wrap_text=True)

        for c_idx in range(1, 10):
            cell_item = ws_formula.cell(row=row_f, column=c_idx)
            cell_item.border = cell_border
            if cur_fill:
                cell_item.fill = cur_fill

        row_f += 1

    # ---------------------------------------------------------
    # AUTO-FIT COLUMN WIDTHS FOR ALL SHEETS
    # ---------------------------------------------------------
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore title rows from calculation
                if cell.row in (2, 3):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    # Output to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
