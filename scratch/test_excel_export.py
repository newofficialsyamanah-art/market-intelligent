import sys
import os
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__) + "/.."))

from app import create_app
from app.models import User
from app.services.market_size_excel import generate_market_size_excel_bytes
import openpyxl

app = create_app()

with app.app_context():
    # 1. Test generate_market_size_excel_bytes directly
    print("Testing generate_market_size_excel_bytes()...")
    buf = generate_market_size_excel_bytes()
    assert buf is not None
    buf_size = len(buf.getvalue())
    print(f"Generated Excel size: {buf_size:,} bytes")
    assert buf_size > 5000, "Excel file is suspiciously small"

    # 2. Inspect generated sheets with openpyxl
    wb = openpyxl.load_workbook(buf)
    sheet_names = wb.sheetnames
    print("Worksheet names:", sheet_names)
    expected_sheets = [
        "Ringkasan Pasar & KPIs",
        "Top 20 Jersey Sublimasi",
        "Top 20 Jaket & Outerwear",
        "Top 20 Seragam Kerja",
        "Top 20 Brand Nasional",
        "Rumus Omzet Perusahaan"
    ]
    for s in expected_sheets:
        assert s in sheet_names, f"Missing sheet: {s}"
        ws = wb[s]
        print(f"  Sheet '{s}': max_row={ws.max_row}, max_column={ws.max_column}")
        assert ws.max_row >= 5, f"Sheet '{s}' has no content"

    # 3. Test via Flask test_client
    print("\nTesting HTTP endpoint /market-analysis/market-size/export-excel...")
    user = User.query.filter_by(role="business_analyst").first()
    client = app.test_client()
    with client.session_transaction() as sess:
        sess['_user_id'] = str(user.id)
        sess['_fresh'] = True

    resp = client.get('/market-analysis/market-size/export-excel')
    print(f"HTTP Status: {resp.status_code}")
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
    print(f"Content-Type: {resp.content_type}")
    assert "spreadsheetml.sheet" in resp.content_type
    print(f"Content-Disposition: {resp.headers.get('Content-Disposition')}")
    assert "Laporan_Market_Size_Share_Syamanah_2026.xlsx" in resp.headers.get("Content-Disposition", "")
    print(f"Downloaded bytes: {len(resp.data):,}")

    print("\nALL EXCEL EXPORT & VALIDATION TESTS PASSED SUCCESSFULLY!")
